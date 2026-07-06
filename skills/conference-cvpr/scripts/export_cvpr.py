#!/usr/bin/env python3
"""Export normalized CVPR records to SQLite, Excel, Markdown, and JSON."""

from __future__ import annotations

import argparse
import json
import logging
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from _cvpr_common import backup_existing, configure_logging, default_export_dir, default_normalized_path, load_json_records, write_json_with_backup


EXPORT_FIELDS = [
    "paper_id",
    "title",
    "authors",
    "authors_text",
    "year",
    "conference",
    "field",
    "source",
    "abstract",
    "pdf_url",
    "paper_page_url",
    "supplementary_url",
    "code_url",
    "project_url",
    "doi",
    "citation_count",
    "openalex_id",
    "semantic_scholar_id",
    "dblp_key",
]


EXCEL_FIELDS = [
    "title",
    "authors",
    "year",
    "conference",
    "abstract",
    "pdf_url",
    "paper_page_url",
    "code_url",
    "project_url",
    "doi",
    "citation_count",
]


def value_for_export(record: dict[str, Any], field: str) -> Any:
    value = record.get(field)
    if field == "authors":
        if isinstance(value, list):
            return "; ".join(str(author) for author in value)
        return value or ""
    if isinstance(value, list):
        return json.dumps(value, ensure_ascii=False)
    return value


def export_sqlite(records: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    backup_existing(path)
    if path.exists():
        path.unlink()
    columns = ", ".join(f"{field} TEXT" for field in EXPORT_FIELDS)
    with sqlite3.connect(path) as conn:
        conn.execute(f"CREATE TABLE papers ({columns})")
        placeholders = ", ".join("?" for _ in EXPORT_FIELDS)
        conn.executemany(
            f"INSERT INTO papers ({', '.join(EXPORT_FIELDS)}) VALUES ({placeholders})",
            [[value_for_export(record, field) for field in EXPORT_FIELDS] for record in records],
        )
        conn.commit()


def export_excel(records: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    backup_existing(path)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "papers"
    worksheet.append(EXCEL_FIELDS)
    for record in records:
        worksheet.append([value_for_export(record, field) for field in EXCEL_FIELDS])
    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 80)
    workbook.save(path)


def escape_markdown(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("\n", " ").replace("|", "\\|")


def export_markdown(records: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    backup_existing(path)
    headers = ["paper_id", "title", "authors", "pdf_url", "paper_page_url"]
    lines = [
        "# CVPR Papers",
        "",
        f"Total papers: {len(records)}",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for record in records:
        row = [escape_markdown(value_for_export(record, field)) for field in headers]
        lines.append("| " + " | ".join(row) + " |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def export_records(input_file: Path, output_dir: Path, year: int) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    records = load_json_records(input_file)
    paths = {
        "sqlite": output_dir / f"cvpr_{year}_papers.sqlite",
        "excel": output_dir / f"cvpr_{year}_papers.xlsx",
        "markdown": output_dir / f"cvpr_{year}_papers.md",
        "json": output_dir / f"cvpr_{year}_papers.json",
    }
    export_sqlite(records, paths["sqlite"])
    export_excel(records, paths["excel"])
    export_markdown(records, paths["markdown"])
    write_json_with_backup(paths["json"], records)
    return paths


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Export normalized CVPR papers to SQLite, Excel, Markdown, and JSON.")
    parser.add_argument("--year", type=int, required=True, help="CVPR year, for example 2026.")
    parser.add_argument("--input-file", help="Normalized JSON file. Defaults to data/normalized/.../cvpr/YEAR/cvpr_YEAR_normalized.json.")
    parser.add_argument("--output-dir", help="Export directory. Defaults to outputs/computer_vision/cvpr/YEAR/.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    configure_logging("export_cvpr", args.year)
    input_file = Path(args.input_file).expanduser().resolve() if args.input_file else default_normalized_path(args.year)
    output_dir = Path(args.output_dir).expanduser().resolve() if args.output_dir else default_export_dir(args.year)
    paths = export_records(input_file, output_dir, args.year)
    for name, path in paths.items():
        logging.info("Wrote %s export to %s", name, path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
